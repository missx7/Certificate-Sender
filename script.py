'''Certificate Sender
Note:
    before you gonna the program you have to :
    1- make sure to replace the ourtemplate.xlsx e-mails
    2- in the send_certificate function that you replace the e-mail py your persnol email and your personal password
    3- maybe it required your permission to send the email via gmail so check your security permission in your gmail account

    thank you!
'''

import re
import xlrd
import smtplib
from docx import Document
from email import encoders
from datetime import datetime
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart


def docx_replace(regex, replace):
    doc = Document('cLast.docx') #you can replace it by your own template
    for p in doc.paragraphs:
        if regex.search(p.text):
            inline = p.runs
            for i in range(len(inline)):
                if regex.search(inline[i].text):
                    text = regex.sub(replace, inline[i].text, count=1)
                    inline[i].text = text
                    doc.save('new.docx')
    return

def extract_xldr():
    wb = xlrd.open_workbook("ourtemplate.xlsx") #you can replace it by your own template
    sheet = wb.sheet_by_index(0)
    new = [sheet.cell(0,cols).value for cols in range(sheet.ncols)]
    email_row = new.index('email')
    name_row  = new.index('Name  ')
    emails    = [sheet.cell(cols+1,email_row).value for cols in range(sheet.ncols+1)]
    names     = [sheet.cell(cols+1,name_row).value for cols in range(sheet.ncols+1)]
    return names,emails


def send_certificate():
    now = datetime.now()
    docx_replace(re.compile(r"Date"), now.strftime("%Y-%m-%d"))
    for a,b in zip(extract_xldr()[0],extract_xldr()[1]):
        docx_replace(re.compile(r"Student NAME"),a)

        #send e-mail

        msg = MIMEMultipart()
        msg['from'] = 'example@gmail.com' #change it to your email
        msg['To'] = b
        msg['Subject'] = 'ta-Da your certificate is ready!'
        password = 'password' #change it to your email password
        part = MIMEBase('application', "octet-stream")
        part.set_payload(open("new.docx", "rb").read())
        encoders.encode_base64(part)
        part.add_header('Content-Disposition', 'attachment; filename="new.docx"')
        msg.attach(part)
        body = 'you made it '
        msg.attach(MIMEText(body, 'html'))
        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(msg['from'], password)
        server.sendmail(msg['from'], msg['To'], msg.as_string())
        server.quit()
    return 'has been sent successfully'

print(send_certificate())